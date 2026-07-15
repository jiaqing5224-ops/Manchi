"""Meeting report service — scan weekly meeting emails and export to Excel."""

import json
import logging
import os
import re
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional

from sqlalchemy.orm import Session

from app.services.llm.client import chat_complete
from app.services.llm.prompts import MEETING_EXTRACT_PROMPT
from app.services.outlook.mail_handler import OutlookMail, scan_inbox_range

logger = logging.getLogger(__name__)


@dataclass
class MeetingInfo:
    """A single meeting extracted from an email."""

    subject: str
    start_time: Optional[datetime]
    end_time: Optional[datetime]
    duration_hours: Optional[float]


def get_this_week_range() -> tuple[datetime, datetime]:
    """Return [Monday 00:00, Saturday 00:00) for the current week."""
    today = datetime.now()
    monday = (today - timedelta(days=today.weekday())).replace(
        hour=0, minute=0, second=0, microsecond=0
    )
    saturday = monday + timedelta(days=5)  # 周六 00:00 == 周五结束
    return monday, saturday


def get_exports_dir() -> str:
    """Return the dedicated directory for Manchi export artifacts.

    All generated artifacts (meeting Excel, etc.) live here so users can find
    them in one place instead of cluttering the Desktop.
    """
    exports = Path(os.environ.get("USERPROFILE", str(Path.home()))) / "Documents" / "Manchi" / "exports"
    exports.mkdir(parents=True, exist_ok=True)
    return str(exports)


def generate_meeting_excel(
    db: Session, output_path: Optional[str] = None
) -> dict:
    """Scan this week's mails, extract meetings, write an Excel to the exports dir.

    Returns a summary dict with path, count, scanned, week range.
    """
    start, end = get_this_week_range()
    mails = scan_inbox_range(start, end)
    meetings = _batch_extract_meetings(mails)

    filename = f"会议报告_周{start:%m%d}-{end:%m%d}.xlsx"
    path = output_path or _exports_path(filename)
    _write_excel(meetings, path)

    return {
        "path": path,
        "count": len(meetings),
        "scanned": len(mails),
        "week_start": start.isoformat(),
        "week_end": end.isoformat(),
    }


def _batch_extract_meetings(mails: list[OutlookMail]) -> list[MeetingInfo]:
    """Extract meetings from mails in batches of 10 via LLM."""
    if not mails:
        return []

    meetings: list[MeetingInfo] = []
    batch_size = 10

    for i in range(0, len(mails), batch_size):
        batch = mails[i : i + batch_size]
        mails_text = "\n".join(
            f"{i + j}. 主题: {m.subject} | 正文: {m.full_body[:1500]}"
            for j, m in enumerate(batch)
        )
        prompt = MEETING_EXTRACT_PROMPT.format(count=len(batch), mails=mails_text)
        try:
            result_text = chat_complete([], prompt, temperature=0.0)
            for item in _parse_json_array(result_text):
                if not item.get("is_meeting"):
                    continue
                meetings.append(_parse_meeting(item))
        except Exception as e:
            logger.error(f"batch extract failed at offset {i}: {e}")
            continue

    return meetings


def _write_excel(meetings: list[MeetingInfo], path: str) -> None:
    """Write meetings to an Excel file with styled headers."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "本周会议"

    headers = ["会议主题", "时长(小时)", "开始时间", "结束时间"]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="4FC3F7", end_color="4FC3F7", fill_type="solid"
    )
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for m in meetings:
        ws.append(
            [
                m.subject,
                m.duration_hours if m.duration_hours is not None else "",
                m.start_time.strftime("%Y-%m-%d %H:%M") if m.start_time else "",
                m.end_time.strftime("%Y-%m-%d %H:%M") if m.end_time else "",
            ]
        )

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20

    wb.save(path)


def _exports_path(filename: str) -> str:
    """Return the full path for an artifact inside the Manchi exports dir."""
    return str(Path(get_exports_dir()) / filename)


def _parse_meeting(item: dict) -> MeetingInfo:
    return MeetingInfo(
        subject=str(item.get("subject", "")),
        start_time=_parse_datetime(item.get("start_time")),
        end_time=_parse_datetime(item.get("end_time")),
        duration_hours=_parse_float(item.get("duration_hours")),
    )


def _parse_datetime(val: object) -> Optional[datetime]:
    if not isinstance(val, str):
        return None
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M"):
        try:
            return datetime.strptime(val, fmt)
        except ValueError:
            continue
    return None


def _parse_float(val: object) -> Optional[float]:
    if isinstance(val, (int, float)):
        return float(val)
    return None


def _parse_json_array(text: str) -> list[dict]:
    """Parse a JSON array from LLM text, tolerating markdown fences."""
    text = text.strip()
    md_match = re.search(r"```(?:json)?\s*\n?(.*?)\n?```", text, re.DOTALL)
    if md_match:
        text = md_match.group(1).strip()
    arr_match = re.search(r"\[.*\]", text, re.DOTALL)
    if arr_match:
        text = arr_match.group(0)
    return json.loads(text)
