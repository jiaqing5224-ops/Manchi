"""File-capability component template (category: file).

Reads/writes Excel (or other files) and returns structured data. When the
upstream input is already `list[dict]` (input_requirement excel/csv/table),
iterate it directly — do NOT re-open the file. Only import openpyxl lazily
inside run() and declare it in manifest.requires.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any


def run(params: dict, context: dict) -> Any:
    """params: 用户配置；context: 运行时上下文（见 spec.md）。

    本模板演示：从 params.file_path 读取一个 Excel 工作表，返回 list[dict]。
    若上游已通过 input_requirement=excel 传入 list[dict]，直接用 context["input_data"]。
    """
    # 情况 A：上游已是解析后的行字典列表
    data = context.get("input_data")
    if isinstance(data, list) and params.get("from_upstream"):
        return _transform(data, params)

    # 情况 B：按路径读取（standalone / input_requirement=file）
    file_path = params.get("file_path") or context.get("source_file_path") or ""
    if not file_path:
        raise ValueError("未提供 file_path，且上游也未传入数据")
    if not Path(file_path).exists():
        raise FileNotFoundError(f"文件不存在: {file_path}")

    from openpyxl import load_workbook  # 惰性导入重型依赖

    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheet_name = params.get("sheet_name") or wb.sheetnames[0]
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    records = []
    for r in rows[1:]:
        if all(c is None for c in r):
            continue
        records.append({header[i]: ("" if r[i] is None else r[i]) for i in range(len(header))})
    return _transform(records, params)


def _transform(records: list[dict], params: dict) -> list[dict]:
    """在这里做你的转换逻辑（过滤/映射/聚合）。"""
    return records


if __name__ == "__main__":
    # 自测：用已解析后的 list[dict] 形状，不要去读文件
    sample = [
        {"Date": "2026-07-27", "Project": "A", "Hours": 4},
        {"Date": "2026-07-28", "Project": "B", "Hours": 3.5},
    ]
    ctx = {"input_data": sample, "artifacts_dir": "/tmp"}
    print(run({"from_upstream": True}, ctx))
